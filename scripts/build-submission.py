"""Submission artifacts only. Does not import or change product runtime."""
from pathlib import Path
import sys, textwrap, json, subprocess, wave
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools/submission'))
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = ROOT / 'docs/submission'
OUT.mkdir(parents=True, exist_ok=True)
NAVY, CREAM, GOLD, INK = '#1c3850', '#faf6ed', '#eebd4a', '#172d42'
W, H = 1600, 900
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
FONT = Path('C:/Windows/Fonts')
def font(size, bold=False):
    return ImageFont.truetype(str(FONT / ('arialbd.ttf' if bold else 'arial.ttf')), size)

titles = [
 'One person. Five economic roles.',
 'Muse interprets. Deterministic engines execute.',
 'Financial evidence reaches one ledger.',
 'Rp750,000 with an evidence trail.',
 'Uncertainty stays with its own scope.',
 'The result survives restart and replay.'
]
narration = [
 "Meet Youssef, a synthetic Moroccan expat. He speaks Darija, receives employment pay in pounds, manages household spending, owns a mixed-use villa, runs a cafe and invests in a gym. BiG-M Lite keeps those economic roles separate. The final demonstration consistently uses October twenty twenty-five. These are actual captured product screens with synthetic records, not real customer data.",
 "Muse Spark interprets conversations and documents and proposes structured candidates. BiG-M validates those candidates. A person reviews consequential facts, transactions and legal content. The existing accounting engine posts balanced journals, and one universal tax engine executes approved rules. Official evidence stays attached to the rule. Muse selects grounded facts to explain the result, while deterministic state supplies every amount.",
 "Conversation, documents and external imports converge on the same posting boundary. The journey includes foreign currencies, a shared credit card, household purchases and two installments. A card repayment reduces a liability instead of creating another expense. Salary, cafe turnover, villa costs and gym capital retain their own meaning. Duplicate imports and confirmations cannot create duplicate journals.",
 "Here is the concrete supported result. The synthetic Indonesian company has one hundred fifty million rupiah in monthly gross turnover. The governed half-percent rate produces seven hundred fifty thousand rupiah in tax. Supported deductions against this final gross base are zero. Explicit fixture evidence confirms no withholding or advance tax. The user can inspect facts, the effective period, official quotations, formula and calculation trace.",
 "The company result does not settle Youssef's worldwide personal taxes. Salary treatment, residence, treaties and mixed-use villa deductions still require evidence. Villa operating costs remain recorded, including staff, cleaning, management and maintenance, without inventing a rental allocation. The supported company amount remains visible while those other components stay under review. This separation makes uncertainty specific and useful.",
 "The full product journey records one hundred twenty-nine balanced journals and preserves state across a fresh process, replayed imports and language changes. One controlled live Muse request explains the calculated tax in Darija, with no retry and no financial-state change. The separate twenty-request residence pilot remains frozen and unresolved. BiG-M Lite shows the source of each supported number and states what it still cannot determine."
]
notes = [
 'Synthetic 2025 persona. Actual product captures from the final browser journey.',
 'One Accounting Engine and one Universal Tax Engine. Muse has no approval authority.',
 'Original currency and reviewed fixture FX remain visible. Card repayment is liability settlement.',
 'Official sources: PP 55/2022 articles 56-62,74; UU7/2021 article 4(3)(c). Simulated human review.',
 'Company tax is separate from personal worldwide tax. No invented deductions or allocation.',
 'See MOROCCAN-EXPAT-E2E.md, deterministic-result.json and tax-live-explanation.json for evidence.'
]
all_images=[]
for index,title in enumerate(titles):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(CREAM[1:])
    image=Image.new('RGB',(W,H),CREAM); d=ImageDraw.Draw(image)
    def rect(x,y,w,h,color,radius=0):
        d.rectangle((x,y,x+w,y+h),fill=color)
        sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color[1:]);sh.line.fill.background()
    def text(x,y,w,value,size=34,color=INK,bold=False):
        # Pixel-based wrapping matches rendered bitmap and editable PowerPoint text.
        f=font(size,bold);lines=[]
        for para in value.split('\n'):
            line=''
            for word in para.split():
                candidate=(line+' '+word).strip()
                if d.textlength(candidate,font=f)>w and line: lines.append(line);line=word
                else:line=candidate
            lines.append(line)
        rendered='\n'.join(lines);d.multiline_text((x,y),rendered,font=f,fill=color,spacing=int(size*.3))
        tb=slide.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches(w/100),Inches((len(lines)*size*1.35+20)/100))
        tf=tb.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for n,line in enumerate(lines):
            p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size*.72);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:])
    def screenshot(name,x,y,w):
        file=ROOT/'research/moroccan-expat/screenshots'/name
        pic=Image.open(file).convert('RGB');h=round(pic.height*w/pic.width)
        image.paste(pic.resize((w,h),Image.Resampling.LANCZOS),(x,y))
        slide.shapes.add_picture(str(file),Inches(x/100),Inches(y/100),width=Inches(w/100))
    rect(0,0,W,14,GOLD)
    text(65,55,1460,title,51,bold=True)
    if index==0:
        text(65,165,500,'BiG-M Lite',58,NAVY,True)
        text(65,265,485,'Darija conversation\nGBP employment\nHousehold and installments\nVilla, cafe and gym',33)
        text(65,610,490,'October 2025\nSynthetic demonstration',27)
        screenshot('dashboard.png',600,180,940)
    elif index==1:
        steps=[('Muse Spark','Interpret context\nPropose facts',65),('Human review','Validate evidence\nConfirm scope',575),('Existing engines','Post journals\nCalculate tax',1085)]
        for label,body,x in steps:
            rect(x,235,450,80,NAVY);text(x+20,250,410,label,32,CREAM,True);text(x+20,350,410,body,32)
        text(65,555,1430,'CountryPack evidence and approved rules govern deterministic execution.',39,bold=True)
        text(65,680,1430,'Muse selects grounded explanation facts. Recorded state supplies the numbers.',31)
    elif index==2:
        text(65,180,470,'Documents\nConversation\nExternal statements',39,NAVY,True)
        text(65,400,470,'Reviewed candidates\nOne posting boundary\n129 balanced journals',34)
        screenshot('reports.png',600,180,940)
    elif index==3:
        text(65,185,480,'IDR 150,000,000',43,NAVY,True)
        text(65,250,480,'Monthly gross base',30)
        text(65,365,480,'0.5%',66,NAVY,True)
        text(65,445,480,'PP 55/2022',31)
        text(65,560,480,'IDR 750,000',53,NAVY,True)
        text(65,635,480,'Remaining payable\nZero withholding / credits',29)
        screenshot('tax.png',600,180,940)
    elif index==4:
        rect(65,215,690,75,NAVY);text(85,230,650,'Supported company calculation',31,CREAM,True)
        text(85,330,630,'October 2025\nFinal gross-turnover tax\nIDR 750,000 payable',39)
        rect(835,215,690,75,GOLD);text(855,230,650,'Review required',31,NAVY,True)
        text(855,330,630,'Personal residence and salary\nMixed-use villa tax treatment\nTreaty relief and credits',35)
        text(85,650,1430,'Villa costs remain recorded. No invented rental allocation or worldwide total.',32)
    else:
        text(65,210,650,'129 balanced journals',52,NAVY,True)
        text(65,305,650,'Fresh-process restart\nDuplicate replay\nLanguage invariance',36)
        text(860,210,650,'1 live explanation',52,NAVY,True)
        text(860,305,650,'Darija\nZero retries\nFinancial state unchanged',36)
        text(65,625,1430,'The 20-request residence pilot stays frozen and unresolved.',30)
        text(65,705,1430,'Facts, official evidence and arithmetic stay inspectable.',38,bold=True)
    text(65,845,1460,f'{index+1}/6   BiG-M Lite   Synthetic product demonstration',20)
    slide.notes_slide.notes_text_frame.text=narration[index]+'\n\n'+notes[index]
    file=OUT/f'slide-{index+1:02}.png';image.save(file);all_images.append(image)
prs.save(OUT/'BiG-M-Lite-Architecture.pptx')
contact=Image.new('RGB',(1200,1013),'#d8d8d8')
for i,img in enumerate(all_images):contact.paste(img.resize((600,338)),((i%2)*600,(i//2)*338))
contact.save(OUT/'slides-preview.png')
script='# Three-minute demo narration\n\nEdited walkthrough of actual UI screenshots. Synthetic stock narration; no voice imitation. Each chapter lasts 30 seconds. This is not a continuous live recording.\n\n'
for i,words in enumerate(narration):
    (OUT/f'narration-{i+1:02}.txt').write_text(words,encoding='utf8')
    script+=f'## {i*30:03d}-{(i+1)*30:03d} seconds: {titles[i]}\n\n{words}\n\n'
(OUT/'VIDEO-SCRIPT.md').write_text(script.rstrip()+'\n',encoding='utf8')
print(json.dumps({'slides':6,'deck':'docs/submission/BiG-M-Lite-Architecture.pptx','narrationWords':sum(len(x.split()) for x in narration)}))

if '--video' in sys.argv:
    import imageio_ffmpeg
    ffmpeg=imageio_ffmpeg.get_ffmpeg_exe()
    clips=[]
    for i in range(1,7):
        wav=ROOT/f'data/submission/narration-{i:02}.wav'
        with wave.open(str(wav)) as a: duration=a.getnframes()/a.getframerate()
        tempo=max(1.0,duration/28.8)
        clip=ROOT/f'data/submission/clip-{i:02}.mp4'
        subprocess.run([ffmpeg,'-v','error','-y','-loop','1','-framerate','10','-i',str(OUT/f'slide-{i:02}.png'),'-i',str(wav),'-vf','scale=1600:900,format=yuv420p','-af',f'atempo={tempo},apad','-t','30','-c:v','libx264','-preset','fast','-crf','23','-c:a','aac','-b:a','128k',str(clip)],check=True)
        clips.append(clip)
    listing=ROOT/'data/submission/clips.txt';listing.write_text('\n'.join("file '"+str(p).replace('\\','/')+"'" for p in clips),encoding='utf8')
    video=OUT/'BiG-M-Lite-3-minute-demo.mp4'
    subprocess.run([ffmpeg,'-v','error','-y','-f','concat','-safe','0','-i',str(listing),'-c','copy','-movflags','+faststart',str(video)],check=True)
    meta=imageio_ffmpeg.read_frames(str(video),pix_fmt='rgb24');info=next(meta);meta.close()
    assert 179.8<=info['duration']<=180.2,info
    assert info['size']==(1600,900),info
    (OUT/'video-verification.json').write_text(json.dumps({'durationSeconds':info['duration'],'dimensions':info['size'],'framesPerSecond':info['fps'],'chapters':6,'chapterDurationSeconds':30,'actualProductScreenshots':3,'continuousScreenRecording':False,'narration':'Stock Windows speech synthesis; no voice imitation','videoBytes':video.stat().st_size},indent=2)+'\n')
    print(json.dumps({'video':str(video.relative_to(ROOT)),'duration':info['duration'],'bytes':video.stat().st_size}))
